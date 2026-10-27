"""
Manav Rachna University – Project Comparator
Sponsored Project + Consultancy Project in one Streamlit app
Developed by aghaZ
"""

import streamlit as st
import re
from io import BytesIO

# Optional: PDF & DOCX extraction
try:
    import pdfplumber
    PDF_OK = True
except ImportError:
    PDF_OK = False

try:
    from docx import Document
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

# ─────────────────────────────────────────────────────────────
# PAGE CONFIG
# ─────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="MRU Project Comparator | Developed by aghaZ",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ─────────────────────────────────────────────────────────────
# CUSTOM CSS
# ─────────────────────────────────────────────────────────────
st.markdown("""
<style>
    .main-header {
        background: linear-gradient(135deg, #1a365d, #2b6cb0);
        color: white;
        padding: 1.4rem 1.8rem;
        border-radius: 12px;
        margin-bottom: 1.2rem;
        position: relative;
    }
    .main-header h1 { margin: 0 0 0.3rem 0; font-size: 1.7rem; }
    .main-header p  { margin: 0; opacity: 0.9; font-size: 0.95rem; }
    .developed-by {
        position: absolute;
        top: 1rem;
        right: 1.5rem;
        font-size: 0.85rem;
        opacity: 0.9;
        font-weight: 500;
    }
    .score-box {
        display: flex;
        align-items: center;
        gap: 1.5rem;
        margin: 1rem 0;
    }
    .score-circle {
        width: 100px;
        height: 100px;
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 1.6rem;
        font-weight: 700;
        color: white;
        flex-shrink: 0;
    }
    .section-card {
        padding: 0.75rem 1rem;
        border-radius: 8px;
        margin-bottom: 0.5rem;
        border-left: 5px solid #ccc;
    }
    .present  { background: #f0fff4; border-left-color: #38a169; }
    .partial  { background: #fffaf0; border-left-color: #dd6b20; }
    .missing  { background: #fff5f5; border-left-color: #e53e3e; }
    .badge {
        display: inline-block;
        padding: 2px 10px;
        border-radius: 4px;
        font-size: 0.75rem;
        font-weight: 600;
        text-transform: uppercase;
    }
    .badge-ok      { background: #c6f6d5; color: #22543d; }
    .badge-partial { background: #feebc8; color: #7b341e; }
    .badge-miss    { background: #fed7d7; color: #742a2a; }
    div[data-testid="stTabs"] button { font-size: 1.05rem; font-weight: 600; }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────
# REQUIRED SECTIONS – SPONSORED PROJECT
# ─────────────────────────────────────────────────────────────
SPONSORED_SECTIONS = [
    {
        "id": "title_page",
        "name": "Title Page / Cover",
        "keywords": ["project title", "manav rachna", "principal investigator", "assistant professor", "associate professor", "department of computer science", "sponsored"],
        "weight": 5,
        "description": "Project title, PI name, designation, department, university, funding agency",
    },
    {
        "id": "completion_certificate",
        "name": "Project Completion Certificate",
        "keywords": ["project completion certificate", "to whom it may concern", "work period", "successfully completed", "authorized signatory"],
        "weight": 10,
        "description": "Official certificate from funding agency with dates and signature",
    },
    {
        "id": "completion_report_header",
        "name": "Project Completion Report Header",
        "keywords": ["project completion report", "project title", "company", "pi details", "co pi details"],
        "weight": 8,
        "description": "Structured report with Project Title, Company, PI Details, Co-PI",
    },
    {
        "id": "objectives",
        "name": "Approved Objectives of Research Project",
        "keywords": ["approved objectives", "objective 1", "objective 2", "objective 3", "objectives of research"],
        "weight": 12,
        "description": "Numbered research objectives (usually 3–4 clear objectives)",
    },
    {
        "id": "dates",
        "name": "Date of Start & Date of Completion",
        "keywords": ["date of start", "date of completion", "work period"],
        "weight": 6,
        "description": "Clear start and end dates of the project",
    },
    {
        "id": "amount",
        "name": "Consultancy / Project Amount",
        "keywords": ["consultancy amount", "amount", "rs.", "inr", "including gst", "budget"],
        "weight": 6,
        "description": "Sanctioned amount (with or without GST)",
    },
    {
        "id": "methodology",
        "name": "Methodology",
        "keywords": ["methodology", "data collection", "data preprocessing", "data analysis", "topic modeling", "sentiment analysis", "visualization"],
        "weight": 15,
        "description": "Detailed step-by-step methodology",
    },
    {
        "id": "outcomes",
        "name": "Outcomes of Research Proposal / Work",
        "keywords": ["outcomes of research", "outcomes of the proposal", "research paper", "published", "findings will help", "insights into"],
        "weight": 12,
        "description": "Statement of outcomes + evidence",
    },
    {
        "id": "research_paper_or_deliverable",
        "name": "Research Paper / Deliverable Evidence",
        "keywords": ["research paper", "doi:", "ieee", "conference", "published", "screenshot", "web app", "mobile app", "chatbot"],
        "weight": 10,
        "description": "Mandatory outcome: published paper or screenshot of developed system",
    },
    {
        "id": "proposal_section",
        "name": "Original Project Proposal",
        "keywords": ["project proposal", "table of content", "about mru", "introduction", "research questions", "objective and expected outcome", "budget", "conclusion"],
        "weight": 8,
        "description": "Full proposal with About MRU, Introduction, Research Questions, Methodology, Budget, Conclusion",
    },
    {
        "id": "request_letter",
        "name": "Request / Approval Letters",
        "keywords": ["request for project grant", "sub:-", "dean research", "hod, cst", "approval"],
        "weight": 4,
        "description": "Internal request letter and approval trail",
    },
    {
        "id": "list_of_projects",
        "name": "University Sponsored Projects List (optional but preferred)",
        "keywords": ["sl.no.", "name of the principal investigator", "funding agency", "amount (in lakhs)", "year"],
        "weight": 4,
        "description": "Extract from university list showing the project entry",
    },
]

# ─────────────────────────────────────────────────────────────
# REQUIRED SECTIONS – CONSULTANCY PROJECT
# ─────────────────────────────────────────────────────────────
CONSULTANCY_SECTIONS = [
    {
        "id": "company_request_letter",
        "name": "Company Request Letter for Consultancy",
        "keywords": ["request for consultancy services", "interested in", "submit a proposal", "cost involved", "authorized signatory"],
        "weight": 8,
        "description": "Letter from the client company requesting consultancy services",
    },
    {
        "id": "internal_approval",
        "name": "Internal Approval / Allotment Letter (Dean Research)",
        "keywords": ["approval of consultancy assignment", "recommended that these consultancy", "allotted to following faculty", "dean research", "principal investigator"],
        "weight": 7,
        "description": "Internal note recommending faculty allotment and seeking approval",
    },
    {
        "id": "proposal_cover",
        "name": "Proposal Cover / Title Page",
        "keywords": ["proposal for", "consultancy services", "associate professor", "department of computer science", "manav rachna university"],
        "weight": 5,
        "description": "Formal proposal title page with project title, company name, PI details",
    },
    {
        "id": "about_mru",
        "name": "About MRU Section",
        "keywords": ["about mru", "manav rachna university", "state private university", "haryana state legislature", "naac"],
        "weight": 4,
        "description": "Standard university introduction paragraph",
    },
    {
        "id": "scope_of_work",
        "name": "Scope of Work / Objectives",
        "keywords": ["scope of work", "identify the current software", "research and recommend", "design and implement", "provide user and technical support", "monitor the system after"],
        "weight": 10,
        "description": "Clear list of objectives / scope items",
    },
    {
        "id": "work_schedule",
        "name": "Work Schedule / Phases",
        "keywords": ["work schedule", "phase 1", "phase 2", "phase 3", "phase 4", "phase 5", "research and assessment"],
        "weight": 9,
        "description": "Phased work plan",
    },
    {
        "id": "proposed_budget",
        "name": "Proposed Budget",
        "keywords": ["proposed budget", "travelling expenses", "consultant fees", "university fees", "total"],
        "weight": 8,
        "description": "Budget table (Travelling, Consultant Fees, University Fees, Total)",
    },
    {
        "id": "completion_report_header",
        "name": "Consultancy Completion Report (Header Table)",
        "keywords": ["consultancy project title", "company:", "pi details", "co pi details", "approved objectives of consultancy", "date of start", "date of completion", "consultancy amount"],
        "weight": 12,
        "description": "Structured table with Title, Company, PI, Co-PI, Objectives, Dates, Amount",
    },
    {
        "id": "methodology",
        "name": "Methodology",
        "keywords": ["methodology", "initial assessment", "planning", "implementation", "evaluation", "monitoring the upgraded system"],
        "weight": 10,
        "description": "Step-wise methodology actually followed",
    },
    {
        "id": "achievements",
        "name": "Consultancy Achievements",
        "keywords": ["consultancy achievements", "increased efficiency", "improved user experience", "reduced operating costs", "improved customer satisfaction"],
        "weight": 8,
        "description": "List of concrete achievements / benefits delivered",
    },
    {
        "id": "outcomes",
        "name": "Outcomes of Consultancy Work",
        "keywords": ["outcomes of consultancy", "software system upgradation is important", "enhanced the efficiency", "improved user experience"],
        "weight": 7,
        "description": "Narrative summary of outcomes and value delivered",
    },
    {
        "id": "completion_certificate",
        "name": "Project Completion Certificate (from Client)",
        "keywords": ["project completion certificate", "to whom it may concern", "successfully completed the project", "work period", "authorized signatory"],
        "weight": 10,
        "description": "Official certificate issued by the client company",
    },
    {
        "id": "distribution_proposal",
        "name": "Consultancy Distribution Proposal (Annexure / Internal Form)",
        "keywords": ["consultancy distribution proposal", "for internal use only", "total consultancy fees received", "share of university", "disbursement to pi", "dean, research", "registrar", "vice chancellor"],
        "weight": 7,
        "description": "Internal university form for fee distribution and signatures",
    },
    {
        "id": "request_for_report",
        "name": "Client Letter Requesting Final Report (optional but good)",
        "keywords": ["request for submission of complete consultancy report", "provide us complete report"],
        "weight": 5,
        "description": "Letter from company asking for the final report",
    },
]

# ─────────────────────────────────────────────────────────────
# TEXT EXTRACTION HELPERS
# ─────────────────────────────────────────────────────────────
def extract_text_from_pdf(file_bytes: bytes) -> str:
    if not PDF_OK:
        st.error(
            "**pdfplumber is not installed on the server.**\n\n"
            "1. Make sure your GitHub repo has a file named `requirements.txt` containing:\n"
            "```\nstreamlit>=1.28.0\npdfplumber>=0.10.3\npython-docx>=1.1.0\n```\n"
            "2. In Streamlit Cloud → your app → **Manage app** → **Reboot app** "
            "(or click **Redeploy**)."
        )
        return ""
    text_parts = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    if not DOCX_OK:
        st.error(
            "**python-docx is not installed on the server.**\n\n"
            "1. Make sure your GitHub repo has a file named `requirements.txt` containing:\n"
            "```\nstreamlit>=1.28.0\npdfplumber>=0.10.3\npython-docx>=1.1.0\n```\n"
            "2. In Streamlit Cloud → your app → **Manage app** → **Reboot app** "
            "(or click **Redeploy**)."
        )
        return ""
    doc = Document(BytesIO(file_bytes))
    parts = []
    # Paragraphs
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            parts.append(p.text.strip())
    # Tables – many consultancy reports keep key content inside tables
    for table in doc.tables:
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = (cell.text or "").strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                parts.append(" | ".join(row_texts))
    # Headers & footers
    for section in doc.sections:
        for hf in (section.header, section.footer):
            if hf is None:
                continue
            for p in hf.paragraphs:
                if p.text and p.text.strip():
                    parts.append(p.text.strip())
    return "\n".join(parts)


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    # Always read from the start (Streamlit can leave the pointer at the end)
    uploaded_file.seek(0)
    data = uploaded_file.read()
    if not data:
        st.error("Uploaded file appears empty.")
        return ""
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        return extract_text_from_docx(data)
    else:
        st.error("Please upload a PDF or DOCX file.")
        return ""


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip()


# ─────────────────────────────────────────────────────────────
# COMPARISON ENGINE
# ─────────────────────────────────────────────────────────────
def compare(text: str, sections: list) -> tuple[int, list]:
    text_n = normalize(text)
    results = []
    total_w = 0
    earned_w = 0

    for sec in sections:
        total_w += sec["weight"]
        hits = sum(1 for kw in sec["keywords"] if kw.lower() in text_n)
        ratio = hits / len(sec["keywords"]) if sec["keywords"] else 0

        if ratio >= 0.40:
            status, points = "present", sec["weight"]
            remark = f"Found {hits}/{len(sec['keywords'])} key indicators. Section appears present."
        elif ratio >= 0.18:
            status, points = "partial", round(sec["weight"] * 0.5)
            remark = f"Only {hits}/{len(sec['keywords'])} indicators found. Section may be incomplete."
        else:
            status, points = "missing", 0
            remark = f"Almost no indicators found ({hits}/{len(sec['keywords'])}). Section is missing or not identifiable."

        earned_w += points
        results.append({**sec, "status": status, "remark": remark, "hits": hits, "points": points})

    score = round((earned_w / total_w) * 100) if total_w else 0
    return score, results


def score_color(score: int) -> str:
    if score >= 85:
        return "#276749"
    if score >= 65:
        return "#2b6cb0"
    if score >= 40:
        return "#c05621"
    return "#c53030"


def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent – Almost complete dossier"
    if score >= 65:
        return "Good – Minor gaps remaining"
    if score >= 40:
        return "Needs Improvement – Several sections missing"
    return "Incomplete – Major sections missing"


# ─────────────────────────────────────────────────────────────
# BUILD DOCX REPORT
# ─────────────────────────────────────────────────────────────
def build_docx_report(score: int, results: list, file_name: str, sample_name: str, project_type: str) -> bytes:
    """Generate a Word report of the comparison and return as bytes."""
    from datetime import datetime
    doc = Document()
    doc.add_heading(f"{project_type} – Comparison Report", level=0)
    doc.add_paragraph(f"Generated by: aghaZ · Manav Rachna University Project Comparator")
    doc.add_paragraph(f"Date: {datetime.now().strftime('%d %B %Y, %I:%M %p')}")
    doc.add_paragraph(f"Uploaded file: {file_name}")
    doc.add_paragraph(f"Reference sample: {sample_name}")
    doc.add_paragraph(f"Overall Compliance Score: {score}%  —  {score_label(score)}")
    doc.add_paragraph("")

    doc.add_heading("Section-by-Section Checklist", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Section / Item"
    hdr[1].text = "Status"
    hdr[2].text = "Finding / Remark"
    for r in results:
        row = table.add_row().cells
        row[0].text = r["name"]
        row[1].text = r["status"].upper()
        row[2].text = r["remark"]

    doc.add_paragraph("")
    doc.add_heading("Missing / Incomplete Items", level=1)
    missing = [r for r in results if r["status"] != "present"]
    if not missing:
        doc.add_paragraph("All required sections appear to be present.")
    else:
        for m in missing:
            doc.add_paragraph(f"• {m['name']}: {m['remark']}", style="List Bullet")
            doc.add_paragraph(f"  Expected: {m['description']}")

    doc.add_paragraph("")
    doc.add_heading("Recommendations", level=1)
    recs = [r for r in results if r["status"] != "present"]
    if not recs:
        doc.add_paragraph("No major recommendations. The file closely matches the sample structure.")
    else:
        for r in recs:
            doc.add_paragraph(f"• Add / complete {r['name']}: {r['description']}", style="List Bullet")

    doc.add_paragraph("")
    doc.add_paragraph("— End of Report —")
    doc.add_paragraph("Developed by aghaZ · Manav Rachna University")

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
# RENDER RESULTS
# ─────────────────────────────────────────────────────────────
def render_results(score: int, results: list, file_name: str, sample_name: str, project_type: str = "Project"):
    color = score_color(score)
    label = score_label(score)

    st.markdown(f"""
    <div class="score-box">
        <div class="score-circle" style="background:{color};">{score}%</div>
        <div>
            <h3 style="margin:0 0 0.3rem 0;">{label}</h3>
            <p style="margin:0; color:#4a5568;">Compared against sample: <em>{sample_name}</em><br>
            File: <strong>{file_name}</strong></p>
        </div>
    </div>
    """, unsafe_allow_html=True)

    st.progress(score / 100)

    st.subheader("Section-by-Section Checklist")
    for r in results:
        cls = r["status"]
        badge_cls = {"present": "badge-ok", "partial": "badge-partial", "missing": "badge-miss"}[cls]
        badge_txt = cls.upper()
        st.markdown(f"""
        <div class="section-card {cls}">
            <strong>{r['name']}</strong>
            <span class="badge {badge_cls}" style="margin-left:8px;">{badge_txt}</span>
            <br><small style="color:#718096;">{r['description']}</small>
            <br><span style="font-size:0.9rem;">{r['remark']}</span>
        </div>
        """, unsafe_allow_html=True)

    missing = [r for r in results if r["status"] != "present"]
    st.subheader("Missing / Incomplete Items")
    if not missing:
        st.success("All required sections appear to be present. Great work!")
    else:
        for m in missing:
            st.warning(f"**{m['name']}** — {m['remark']}\n\n*Expected:* {m['description']}")

    st.subheader("Recommendations")
    recs = []
    for r in results:
        if r["status"] == "present":
            continue
        recs.append(f"• Add / complete **{r['name']}**: {r['description']}")
    if recs:
        st.markdown("\n".join(recs))
    else:
        st.success("No major recommendations. The uploaded file closely matches the sample structure.")

    # Download report as DOCX
    if DOCX_OK:
        report_bytes = build_docx_report(score, results, file_name, sample_name, project_type)
        safe_name = re.sub(r"[^\w\-_.]", "_", file_name.rsplit(".", 1)[0])[:60]
        st.download_button(
            label="📥 Download Report (DOCX)",
            data=report_bytes,
            file_name=f"Comparison_Report_{safe_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key=f"dl_{file_name}_{score}",
            use_container_width=True,
        )
    else:
        st.caption("Install python-docx to enable DOCX report download.")


# ─────────────────────────────────────────────────────────────
# MAIN UI
# ─────────────────────────────────────────────────────────────
st.markdown("""
<div class="main-header">
    <div class="developed-by">Developed by aghaZ</div>
    <h1>MRU Project Comparator</h1>
    <p>Manav Rachna University · Compare Sponsored &amp; Consultancy Project files against approved samples</p>
</div>
""", unsafe_allow_html=True)

tab1, tab2 = st.tabs(["📑 Sponsored Project Comparator", "💼 Consultancy Project Comparator"])

MAX_FILES = 10  # maximum files allowed at once

# ─── TAB 1: SPONSORED ───────────────────────────────────────
with tab1:
    st.info("**Reference sample already loaded:** “Covid Data Visualization on Tweet Data” (Dr. Mamta Arora, Lala Madho Ram Bhagwan Dass Charitable Society, 2022-23)")
    st.caption(f"You can upload **up to {MAX_FILES} files** at once (PDF or DOCX).")

    uploaded_list = st.file_uploader(
        "Upload your Sponsored Project file(s) (PDF or DOCX)",
        type=["pdf", "docx"],
        accept_multiple_files=True,
        key="sponsored_uploader",
    )

    if uploaded_list:
        if len(uploaded_list) > MAX_FILES:
            st.warning(f"Maximum {MAX_FILES} files allowed at once. Only the first {MAX_FILES} will be processed.")
            uploaded_list = uploaded_list[:MAX_FILES]

        st.success(f"**{len(uploaded_list)} file(s)** ready (max {MAX_FILES})")
        for f in uploaded_list:
            st.write(f"• {f.name} ({f.size // 1024} KB)")

        if st.button("🔍 Compare All with Sample", key="btn_sponsored", type="primary", use_container_width=True):
            for idx, uploaded in enumerate(uploaded_list):
                st.markdown(f"---\n### File {idx + 1} of {len(uploaded_list)}: `{uploaded.name}`")
                with st.spinner(f"Extracting & comparing {uploaded.name}…"):
                    text = extract_text(uploaded)
                    if text and len(text.strip()) > 30:
                        st.caption(f"Extracted {len(text)} characters.")
                        score, results = compare(text, SPONSORED_SECTIONS)
                        render_results(
                            score,
                            results,
                            uploaded.name,
                            "Covid Data Visualization on Tweet Data",
                            project_type="Sponsored Project",
                        )
                    elif text is not None:
                        st.error(
                            f"Could not extract enough text from this file "
                            f"(only {len(text.strip())} characters found). "
                            "The file may be image-only / scanned, or protected."
                        )
    else:
        st.markdown("""
        **What this tab checks**
        - Title / Cover page
        - Project Completion Certificate
        - Project Completion Report (Title, Company, PI, Co-PI)
        - Approved Objectives
        - Date of Start & Completion
        - Project Amount
        - Detailed Methodology
        - Outcomes + Research Paper / Deliverable evidence
        - Original Project Proposal
        - Request / Approval letters

        **Upload limit:** up to **10 files** at once.
        """)

# ─── TAB 2: CONSULTANCY ─────────────────────────────────────
with tab2:
    st.info("**Reference sample already loaded:** “Software System Upgradation for ACRO” (Dr. Ranjna Jain, ACRO Engineering Company, 2022-23)")
    st.caption(f"You can upload **up to {MAX_FILES} files** at once (PDF or DOCX).")

    uploaded_list2 = st.file_uploader(
        "Upload your Consultancy Project file(s) (PDF or DOCX)",
        type=["pdf", "docx"],
        accept_multiple_files=True,
        key="consultancy_uploader",
    )

    if uploaded_list2:
        if len(uploaded_list2) > MAX_FILES:
            st.warning(f"Maximum {MAX_FILES} files allowed at once. Only the first {MAX_FILES} will be processed.")
            uploaded_list2 = uploaded_list2[:MAX_FILES]

        st.success(f"**{len(uploaded_list2)} file(s)** ready (max {MAX_FILES})")
        for f in uploaded_list2:
            st.write(f"• {f.name} ({f.size // 1024} KB)")

        if st.button("🔍 Compare All with Sample", key="btn_consultancy", type="primary", use_container_width=True):
            for idx, uploaded2 in enumerate(uploaded_list2):
                st.markdown(f"---\n### File {idx + 1} of {len(uploaded_list2)}: `{uploaded2.name}`")
                with st.spinner(f"Extracting & comparing {uploaded2.name}…"):
                    text = extract_text(uploaded2)
                    if text and len(text.strip()) > 30:
                        st.caption(f"Extracted {len(text)} characters.")
                        score, results = compare(text, CONSULTANCY_SECTIONS)
                        render_results(
                            score,
                            results,
                            uploaded2.name,
                            "Software System Upgradation for ACRO",
                            project_type="Consultancy Project",
                        )
                    elif text is not None:
                        st.error(
                            f"Could not extract enough text from this file "
                            f"(only {len(text.strip())} characters found). "
                            "The file may be image-only / scanned, or protected."
                        )
    else:
        st.markdown("""
        **What this tab checks**
        - Company Request Letter
        - Internal Approval / Allotment (Dean Research)
        - Proposal Cover + About MRU
        - Scope of Work / Objectives
        - Work Schedule (Phases)
        - Proposed Budget table
        - Completion Report header table
        - Methodology actually followed
        - Consultancy Achievements & Outcomes
        - Client Project Completion Certificate
        - Consultancy Distribution Proposal (Annexure)

        **Upload limit:** up to **10 files** at once.
        """)

# Footer
st.markdown("---")
st.caption("Developed by **aghaZ** · Manav Rachna University · Project Comparator Tool")
